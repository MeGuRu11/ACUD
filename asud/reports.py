"""Excel and Word report generation."""

from datetime import datetime

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from openpyxl.styles import Border, Side


class ReportGenerator:
    def __init__(self, config):
        self.config = config

    def export_to_excel(self, df, filepath, selected_columns=None):
        if selected_columns: df = df[selected_columns]
        with pd.ExcelWriter(filepath, engine='openpyxl') as w:
            df.to_excel(w, sheet_name='Отчёт', index=False)
            ws = w.sheets['Отчёт']
            b = Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))
            for r in ws.iter_rows():
                for c in r:
                    if c.value is not None: c.border = b
            for col in ws.columns:
                ml = max((len(str(c.value)) for c in col if c.value), default=0)
                ws.column_dimensions[col[0].column_letter].width = min(ml + 2, 50)

    def _remove_cell_borders(self, cell):
        try:
            tc = cell._element
            tcPr = tc.tcPr or OxmlElement('w:tcPr') or tc.append(OxmlElement('w:tcPr'))
            borders = tcPr.find(qn('w:tcBorders')) or tcPr.append(OxmlElement('w:tcBorders'))
            for n in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
                b = borders.find(qn(f'w:{n}'))
                if b is not None: borders.remove(b)
        except:
            pass

    def export_to_word(self, df, filepath, query_text="", selected_columns=None, template_path=None):
        from docx.enum.section import WD_SECTION, WD_ORIENT
        if selected_columns: df = df[selected_columns]
        doc = Document()
        s = doc.sections[0];
        s.page_height = Cm(29.7);
        s.page_width = Cm(21);
        s.top_margin = Cm(2);
        s.bottom_margin = Cm(2);
        s.left_margin = Cm(3);
        s.right_margin = Cm(1.5)
        ht = doc.add_table(1, 2);
        ht.autofit = False;
        ht.columns[0].width = Cm(4);
        ht.columns[1].width = Cm(13)
        ch = ht.cell(0, 0);
        ch.vertical_alignment = WD_ALIGN_PARAGRAPH.CENTER
        for p in ch.paragraphs: p.clear()
        p1 = ch.paragraphs[0];
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run(
            "МИНИСТЕРСТВО ОБОРОНЫ\nРОССИЙСКОЙ ФЕДЕРАЦИИ\n(МИНОБОРОНЫ РОССИИ)\nВОЕННО-МЕДИЦИНСКАЯ АКАДЕМИЯ\nг. Санкт-Петербург, ул. Академика Лебедева, д.6, 194044")
        r1.font.name = 'Times New Roman';
        r1.font.size = Pt(12);
        r1.bold = True
        for row in ht.rows:
            for cell in row.cells: self._remove_cell_borders(cell)
        doc.add_paragraph().paragraph_format.space_after = Pt(18)
        h = doc.add_paragraph();
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER;
        rh = h.add_run("Справка.");
        rh.font.name = 'Times New Roman';
        rh.font.size = Pt(14);
        rh.bold = True;
        h.paragraph_format.space_after = Pt(12)
        i = doc.add_paragraph();
        i.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY;
        i.paragraph_format.line_spacing = 1.15
        ri = i.add_run(
            f"В ответ на Ваш запрос представляем информацию о диссертационных работах, соответствующих следующим критериям: {query_text}." if query_text else "В ответ на Ваш запрос представляем информацию о диссертационных работах.")
        ri.font.name = 'Times New Roman';
        ri.font.size = Pt(12);
        i.paragraph_format.space_after = Pt(12)
        ar = doc.add_paragraph();
        ar.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY;
        ar.paragraph_format.line_spacing = 1.15;
        ar.paragraph_format.space_after = Pt(18)
        ra = ar.add_run(
            "Сведения о диссертационных работах приведены в Приложении на ____ листе(ах) в алфавитном порядке.")
        ra.font.name = 'Times New Roman';
        ra.font.size = Pt(12)
        doc.add_paragraph().paragraph_format.space_after = Pt(36)
        st = doc.add_table(2, 2);
        st.autofit = False;
        st.columns[0].width = Cm(12);
        st.columns[1].width = Cm(5)
        st.cell(0, 0).text = "Начальник отдела по работе с диссертационными советами"
        st.cell(1, 0).text = "врач-методист"
        sn = st.cell(1, 1);
        sn.text = "Стеганцев И.В."
        for p in sn.paragraphs: p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for row in st.rows:
            for cell in row.cells: self._remove_cell_borders(cell)
        doc.add_paragraph().paragraph_format.space_after = Pt(12)
        dp = doc.add_paragraph();
        dp.alignment = WD_ALIGN_PARAGRAPH.LEFT;
        dr = dp.add_run(datetime.now().strftime("%d.%m.%Y"));
        dr.font.name = 'Times New Roman';
        dr.font.size = Pt(11)
        if not df.empty:
            doc.add_page_break()
            ns = doc.add_section(WD_SECTION.NEW_PAGE);
            ns.orientation = WD_ORIENT.LANDSCAPE;
            ns.page_width = Cm(29.7);
            ns.page_height = Cm(21);
            ns.top_margin = Cm(1.5);
            ns.bottom_margin = Cm(1.5);
            ns.left_margin = Cm(1.5);
            ns.right_margin = Cm(1.5)
            ah = doc.add_paragraph();
            ah.alignment = WD_ALIGN_PARAGRAPH.CENTER;
            rht = ah.add_run("ПРИЛОЖЕНИЕ");
            rht.font.name = 'Times New Roman';
            rht.font.size = Pt(14);
            rht.bold = True;
            ah.paragraph_format.space_after = Pt(6)
            ash = doc.add_paragraph();
            ash.alignment = WD_ALIGN_PARAGRAPH.CENTER;
            rs = ash.add_run("Сведения о диссертационных работах");
            rs.font.name = 'Times New Roman';
            rs.font.size = Pt(12);
            ash.paragraph_format.space_after = Pt(18)
            dc = [c for c in df.columns if c != '_original_index']
            tbl = doc.add_table(1, len(dc));
            tbl.style = 'Table Grid';
            tbl.autofit = False
            cw = {'ФИО': Cm(4), 'Диссертационный совет': Cm(3.5), 'Название диссертации': Cm(7),
                  'Дата защиты диссертации': Cm(2.5), 'Специальность': Cm(3),
                  '1 Научный руководитель (консультант)': Cm(4), '2 Научный руководитель (консультант)': Cm(4),
                  'Год защиты': Cm(1.5), 'Искомая степень': Cm(3.5), 'Информация о лишении степени': Cm(4),
                  'Примечания': Cm(4)}
            for i, c in enumerate(dc):
                tbl.rows[0].cells[i].text = c
                for p in tbl.rows[0].cells[
                    i].paragraphs: p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(2)
                for r in p.runs: r.font.name = 'Times New Roman'; r.font.size = Pt(9); r.bold = True
                tbl.rows[0].cells[i].width = cw.get(c, Cm(3))
            for _, row in df.iterrows():
                rc = tbl.add_row().cells
                for i, c in enumerate(dc):
                    v = row[c];
                    t = str(v) if pd.notna(v) and str(v).strip() != '' else '–'
                    rc[i].text = t
                    for p in rc[
                        i].paragraphs: p.alignment = WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.space_after = Pt(2)
                    for r in p.runs: r.font.name = 'Times New Roman'; r.font.size = Pt(9)
                    if c in ['Название диссертации', 'Примечания', 'Информация о лишении степени']: rc[i].paragraphs[
                        0].paragraph_format.word_wrap = True
            fn = doc.add_paragraph();
            fn.alignment = WD_ALIGN_PARAGRAPH.LEFT;
            fn.paragraph_format.space_after = Pt(6);
            rf = fn.add_run(f"Всего записей: {len(df)}");
            rf.font.name = 'Times New Roman';
            rf.font.size = Pt(10);
            rf.italic = True
        doc.save(filepath)
