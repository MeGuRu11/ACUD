import pytest

import pandas as pd
from openpyxl import load_workbook

from asud.reports import ReportGenerator


def test_export_to_excel_rejects_empty_column_selection(tmp_path):
    generator = ReportGenerator({})
    df = pd.DataFrame([{"ФИО": "Иванов И.И.", "Год защиты": 2024}])

    with pytest.raises(ValueError, match="Выберите хотя бы одну колонку"):
        generator.export_to_excel(df, tmp_path / "report.xlsx", selected_columns=[])


def test_export_to_excel_uses_readable_domain_column_widths(tmp_path):
    generator = ReportGenerator({})
    df = pd.DataFrame(
        [
            {
                "Год защиты": 2024,
                "ФИО": "Иванов И.И.",
                "Название диссертации": "Короткое название",
                "Диссертационный совет": "21.2.050.01",
                "Примечания": "Нет",
            }
        ]
    )
    output = tmp_path / "report.xlsx"

    generator.export_to_excel(df, output)

    worksheet = load_workbook(output).active
    widths = {
        cell.value: worksheet.column_dimensions[cell.column_letter].width
        for cell in worksheet[1]
    }
    assert widths["Год защиты"] >= 12
    assert widths["ФИО"] >= 24
    assert widths["Название диссертации"] >= 42
    assert widths["Диссертационный совет"] >= 22
    assert widths["Примечания"] >= 24


def test_export_to_excel_hides_empty_placeholders(tmp_path):
    generator = ReportGenerator({})
    df = pd.DataFrame(
        [
            {
                "ФИО": "Иванов И.И.",
                "Примечания": "nan",
                "Информация о лишении степени": "None",
                "Искомая степень": pd.NA,
            }
        ]
    )
    output = tmp_path / "report.xlsx"

    generator.export_to_excel(df, output)

    worksheet = load_workbook(output).active
    row_values = [cell.value for cell in worksheet[2]]
    assert "nan" not in row_values
    assert "None" not in row_values
    assert "<NA>" not in row_values


def test_word_column_widths_fit_available_page_width():
    columns = [
        "ФИО",
        "Название диссертации",
        "Диссертационный совет",
        "Дата защиты диссертации",
        "Специальность",
        "Искомая степень",
        "Информация о лишении степени",
        "Примечания",
    ]

    widths = ReportGenerator.calculate_word_column_widths(columns, total_width_cm=26.7)

    assert list(widths) == columns
    assert round(sum(widths.values()), 2) <= 26.7
    assert all(width > 0 for width in widths.values())


def test_export_to_word_creates_document_with_selected_columns(tmp_path):
    generator = ReportGenerator({})
    df = pd.DataFrame(
        [
            {
                "ФИО": "Иванов И.И.",
                "Название диссертации": "Исследование",
                "Год защиты": 2024,
            }
        ]
    )
    output = tmp_path / "report.docx"

    generator.export_to_word(df, output, selected_columns=["ФИО", "Год защиты"])

    assert output.exists()
    assert output.stat().st_size > 0
